"""
ወደ ኤክሴልና ፒዲኤፍ መላክ — Excel and PDF export.

Division of labour, deliberately:
  * Excel is DATA — one flat table, real numbers, no merged cells, so HR
    can sort, filter and pivot it or paste it into a payroll sheet.
  * PDF is PRESENTATION — paginated, headed, with a signature and stamp
    block, because anything leaving the HR office gets signed.

The Ethiopic font is registered explicitly and its presence is asserted at
import time. ReportLab renders unregistered Ethiopic glyphs as empty boxes
WITHOUT raising an error, so a missing font produces a silently blank
Amharic report. That failure mode is why this module refuses to start
rather than warn.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from ..core.ethiopian import EthiopianDate
from ..i18n.am import t

FONT_DIR = Path(__file__).resolve().parents[2] / "assets" / "fonts"
FONT_REGULAR = "SAMSAmharic"
FONT_BOLD = "SAMSAmharic-Bold"
FONT_FILE_REGULAR = "SAMSAmharic-Regular.ttf"
FONT_FILE_BOLD = "SAMSAmharic-Bold.ttf"

# Everything a report actually prints: Amharic, digits, and the punctuation
# that appears in hours, percentages and the "ተ.ቁ" column header. Checked at
# startup because a font missing only DIGITS renders names perfectly and
# every number as blank space, with no error raised anywhere.
REQUIRED_CHARS = "0123456789.,%-:/()የሰራተኛስምሰዓትጳጉሜ"

_SEARCH_PATHS = [
    FONT_DIR,
    Path("/usr/share/fonts/truetype/noto"),
    Path("C:/Windows/Fonts"),
]


class FontMissing(RuntimeError):
    """The Ethiopic font is not available, so Amharic cannot be rendered."""


def _find(filename: str) -> Path | None:
    for base in _SEARCH_PATHS:
        p = base / filename
        if p.exists():
            return p
    return None


def register_fonts() -> None:
    """
    Register the Ethiopic font, or refuse to continue.

    Called at import. Failing loudly here is the whole point: the
    alternative is a report full of empty rectangles that looks fine to
    the code and is useless to the office.
    """
    if FONT_REGULAR in pdfmetrics.getRegisteredFontNames():
        return
    regular = _find(FONT_FILE_REGULAR)
    if regular is None:
        raise FontMissing(
            f"{FONT_FILE_REGULAR} not found. Amharic PDFs cannot be rendered "
            f"without it. Place it in {FONT_DIR}, or rebuild it with "
            "`python tools/build_font.py`."
        )
    pdfmetrics.registerFont(TTFont(FONT_REGULAR, str(regular)))
    bold = _find(FONT_FILE_BOLD)
    pdfmetrics.registerFont(TTFont(FONT_BOLD, str(bold or regular)))

    missing = _missing_glyphs(FONT_REGULAR, REQUIRED_CHARS)
    if missing:
        raise FontMissing(
            f"the registered font cannot render {missing!r}. An Ethiopic-only "
            "font has no digits or Latin punctuation, so every number in every "
            "report would come out blank. Rebuild with "
            "`python tools/build_font.py`."
        )


def _missing_glyphs(font_name: str, sample: str) -> list[str]:
    face = pdfmetrics.getFont(font_name).face
    cmap = getattr(face, "charToGlyph", None)
    if not cmap:
        return []
    return sorted({ch for ch in sample if ch.strip() and ord(ch) not in cmap})


def verify_amharic_coverage(
    sample: str = "የሰራተኞች ተገኝነት መቆጣጠሪያ ስርዓት 1234567890.%",
) -> bool:
    """
    Confirm the registered font can render everything a report contains.

    The sample deliberately includes digits and punctuation, not just
    Amharic. Checking only Amharic letters is how a font that renders every
    name correctly and every HOUR as empty space passes a test suite.
    """
    register_fonts()
    return not _missing_glyphs(FONT_REGULAR, sample)


register_fonts()


# -- Excel ---------------------------------------------------------------


def to_excel(report: dict[str, Any], path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "ሪፖርት"
    ws.sheet_view.rightToLeft = False

    amharic = Font(name="Noto Sans Ethiopic", size=11)
    head_font = Font(name="Noto Sans Ethiopic", size=11, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F4A6D")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([report["title_am"]])
    ws["A1"].font = Font(name="Noto Sans Ethiopic", size=14, bold=True)
    ws.append([t("app.office")])
    ws.append([report.get("period_am", "")])
    if report.get("employee_am"):
        ws.append([f"{t('emp.name')}: {report['employee_am']}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(report["columns"])
    for cell in ws[header_row]:
        cell.font = head_font
        cell.fill = head_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)

    for row in report["rows"]:
        ws.append([row.get(c, "") for c in report["columns"]])

    for r in ws.iter_rows(min_row=header_row + 1):
        for cell in r:
            cell.font = amharic
            cell.border = border

    if report.get("totals"):
        ws.append([])
        ws.append(
            [t("total")]
            + [report["totals"].get(c, "") for c in report["columns"][1:]]
        )
        for cell in ws[ws.max_row]:
            cell.font = Font(name="Noto Sans Ethiopic", size=11, bold=True)

    for i, col in enumerate(report["columns"], start=1):
        width = max(len(str(col)) + 4,
                    *(len(str(r.get(col, ""))) + 3 for r in report["rows"] or [{}]))
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = (
            min(42, max(12, width))
        )
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    wb.save(out)
    return out


# -- PDF -----------------------------------------------------------------


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "AmTitle", parent=base["Title"], fontName=FONT_BOLD,
            fontSize=15, leading=20, alignment=TA_CENTER, spaceAfter=2,
        ),
        "office": ParagraphStyle(
            "AmOffice", parent=base["Normal"], fontName=FONT_REGULAR,
            fontSize=11.5, leading=15, alignment=TA_CENTER, spaceAfter=2,
        ),
        "period": ParagraphStyle(
            "AmPeriod", parent=base["Normal"], fontName=FONT_REGULAR,
            fontSize=10.5, leading=14, alignment=TA_CENTER,
            textColor=colors.HexColor("#444444"), spaceAfter=8,
        ),
        "body": ParagraphStyle(
            "AmBody", parent=base["Normal"], fontName=FONT_REGULAR,
            fontSize=9.5, leading=13,
        ),
        "note": ParagraphStyle(
            "AmNote", parent=base["Normal"], fontName=FONT_REGULAR,
            fontSize=8.5, leading=12, textColor=colors.HexColor("#555555"),
        ),
        "cell": ParagraphStyle(
            "AmCell", fontName=FONT_REGULAR, fontSize=8.5, leading=11,
        ),
        "head": ParagraphStyle(
            "AmHead", fontName=FONT_BOLD, fontSize=8.5, leading=11,
            alignment=TA_CENTER, textColor=colors.white,
        ),
    }


def _footer(canvas, doc) -> None:
    """Page number plus the print date, in Amharic and the Ethiopian calendar."""
    canvas.saveState()
    canvas.setFont(FONT_REGULAR, 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    eth = EthiopianDate.today()
    canvas.drawString(
        15 * mm, 10 * mm,
        f"{t('rep.printed_on')}: {eth.format_am()}",
    )
    canvas.drawRightString(
        doc.pagesize[0] - 15 * mm, 10 * mm, f"ገጽ {doc.page}"
    )
    canvas.restoreState()


def to_pdf(
    report: dict[str, Any], path: str | Path, landscape_mode: bool = True
) -> Path:
    register_fonts()
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    st = _styles()

    doc = SimpleDocTemplate(
        str(out),
        pagesize=landscape(A4) if landscape_mode else A4,
        leftMargin=13 * mm, rightMargin=13 * mm,
        topMargin=14 * mm, bottomMargin=18 * mm,
        title=report["title_am"], author=t("app.office"),
    )

    story: list[Any] = [
        Paragraph(t("app.office"), st["office"]),
        Paragraph(report["title_am"], st["title"]),
        Paragraph(report.get("period_am", ""), st["period"]),
    ]
    if report.get("employee_am"):
        story.append(Paragraph(
            f"{t('emp.name')}: {report['employee_am']} — "
            f"{report.get('position_am', '')}", st["body"]))
        story.append(Spacer(1, 4))

    cols = report["columns"]
    data = [[Paragraph(str(c), st["head"]) for c in cols]]
    for row in report["rows"]:
        data.append([Paragraph(str(row.get(c, "")), st["cell"]) for c in cols])

    if len(data) == 1:
        story.append(Paragraph(t("rep.no_data"), st["body"]))
    else:
        table = Table(data, repeatRows=1, hAlign="CENTER")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F4A6D")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C3D1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F2F5F9")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)

    if report.get("totals"):
        story.append(Spacer(1, 6))
        totals = " · ".join(
            f"{k}: {v}" for k, v in report["totals"].items()
        )
        story.append(Paragraph(f"<b>{t('total')}</b> — {totals}", st["body"]))

    if report.get("summary"):
        story.append(Spacer(1, 6))
        summary = " · ".join(f"{k}: {v}" for k, v in report["summary"].items())
        story.append(Paragraph(summary, st["body"]))

    if report.get("note_am"):
        story.append(Spacer(1, 6))
        story.append(Paragraph(report["note_am"], st["note"]))

    story.append(Spacer(1, 14 * mm))
    story.append(_signature_block(st))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return out


def _signature_block(st: dict[str, ParagraphStyle]) -> Table:
    """
    Space for preparer, approver and the official stamp.

    Any report that leaves the HR office will be signed. A report with
    nowhere to sign gets retyped by hand, and then the system has saved
    nobody any work.
    """
    line = "_______________________"
    data = [
        [
            Paragraph(f"{t('rep.prepared_by')}: {line}", st["body"]),
            Paragraph(f"{t('rep.approved_by')}: {line}", st["body"]),
            Paragraph(t("rep.stamp"), st["body"]),
        ],
        [
            Paragraph(f"{t('rep.signature')}: {line}", st["body"]),
            Paragraph(f"{t('rep.signature')}: {line}", st["body"]),
            Paragraph("", st["body"]),
        ],
    ]
    tbl = Table(data, colWidths=[None, None, 45 * mm])
    tbl.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOX", (2, 0), (2, 1), 0.5, colors.HexColor("#999999")),
    ]))
    return tbl
